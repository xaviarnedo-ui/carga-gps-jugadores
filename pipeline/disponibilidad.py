"""Disponibilidad y Minutos.xlsx, regenerado entero desde los microciclos.

Una hoja por mes (el microciclo va al mes en que empieza), una columna por sesión/partido
con los minutos, coloreada por estado. Sin fórmulas (TOTAL, MÁXIMO y % ya calculados), así
no hace falta recalcular con LibreOffice. Las sesiones Extra no aparecen.
"""
import os
import re
import shutil

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from . import config as C
from . import estados as E
from . import sesion
from .pdf_catapult import dur_a_min
from .xlsx import fecha_hoja, filas_jugadores, microciclos_existentes, num

MESES = ["", "Enero", "Febrero", "Marzo", "Abril", "Mayo", "Junio", "Julio", "Agosto",
         "Septiembre", "Octubre", "Noviembre", "Diciembre"]
MARCA = "Generado automáticamente por el pipeline GPS"
AMARILLO = PatternFill("solid", fgColor=C.AMARILLO[0])
ROJO = PatternFill("solid", fgColor=C.ROJO[0])
BANDA = [PatternFill("solid", fgColor="FFFFFFFF"), PatternFill("solid", fgColor="FFF2F5FA")]
AZUL_HDR = PatternFill("solid", fgColor="FF4472C4")
ORO_HDR = PatternFill("solid", fgColor="FFB8860B")
MC_HDR = PatternFill("solid", fgColor="FF2E4E6E")
GRIS = PatternFill("solid", fgColor=C.GRIS_MEDIA)
BLANCA = Font(color="FFFFFFFF", bold=True)
GRUESA = Side(style="medium", color="FF1F3864")


def _eventos(n, wb, registro):
    """[(key, fecha, es_partido, {dorsal: (minutos, estado)})] de las sesiones/partidos con datos."""
    out = []
    for nombre in wb.sheetnames:
        if not C.SHEET_KEY_RE.match(nombre):
            continue
        ws, key = wb[nombre], nombre[:-4]
        partido = C.is_match_key(key)
        filas, _ = filas_jugadores(ws)
        dur_c = C.MAT_DUR if partido else C.SES_DUR
        pl_c = C.MAT_PL if partido else C.SES_PL
        datos = {}
        for d, r in filas.items():
            real = num(ws.cell(r, pl_c).value) is not None or any(
                num(ws.cell(r, (C.MAT_COL[m] if partido else C.SES_OBJ_COL[m] + 1)).value) is not None
                for m in C.METRICS)
            mins = dur_a_min(ws.cell(r, dur_c).value) if ws.cell(r, dur_c).value else None
            datos[d] = {"real": real, "min": mins or 0,
                        "obj": (not partido) and sesion.tiene_objetivo(ws, r)}
        if not any(v["real"] for v in datos.values()):
            continue                                    # aún no se ha hecho
        reg = registro.get(key, {}).get("estados", {})
        res = {}
        for d, v in datos.items():
            est = reg.get(str(d))
            if est is None:                             # histórico previo al pipeline: se deduce
                if v["real"]:
                    est = C.FULL if (partido or v["obj"]) else C.REHAB
                else:
                    est = None if (partido or v["obj"]) else C.LESION
            res[d] = (v["min"] if v["real"] else 0, est)
        out.append((key, fecha_hoja(ws), partido, res))
    out.sort(key=lambda t: (t[1], t[0]))
    # en los partidos, quien no jugó hereda el rojo si en la sesión anterior estaba lesionado
    prev = {}
    for key, fecha, partido, res in out:
        for d, (mins, est) in list(res.items()):
            if partido and est is None and prev.get(d) == C.LESION:
                res[d] = (mins, C.LESION)
        prev = {d: est for d, (_, est) in res.items()}
    return out


def _etiqueta_mc(n, wb):
    for key, ws, _ in sesion.hojas_sesion(wb):
        tipo, _ = sesion.tipo_y_dia(ws)
        if tipo:
            return f"MC{n} · TIPO {tipo}"
    return f"MC{n}"


def generar(ruta=None, abiertos=None):
    ruta = ruta or C.DISPO_XLSX
    abiertos = abiertos or {}
    registro = E.cargar()["sesiones"]
    meses, plantilla = {}, {}
    for n, r in microciclos_existentes().items():
        wb = abiertos.get(n) or openpyxl.load_workbook(r, data_only=True)
        evs = _eventos(n, wb, registro)
        if not evs:
            continue
        for ws in (wb[x] for x in wb.sheetnames if C.SHEET_KEY_RE.match(x)):
            filas, _ = filas_jugadores(ws)
            for d, fila in filas.items():
                plantilla[d] = (ws.cell(fila, 2).value, ws.cell(fila, 3).value)
        mes = (evs[0][1].year, evs[0][1].month)
        meses.setdefault(mes, []).append((_etiqueta_mc(n, wb), evs))

    if os.path.exists(ruta):
        viejo = openpyxl.load_workbook(ruta, read_only=True)
        a2 = str(viejo.worksheets[0].cell(2, 1).value or "")
        viejo.close()
        if MARCA not in a2:                            # el fichero manual antiguo se archiva
            os.makedirs(C.ARCHIVO_DIR, exist_ok=True)
            shutil.move(ruta, os.path.join(C.ARCHIVO_DIR, "Disponibilidad y Minutos (manual, hasta agosto).xlsx"))

    out = openpyxl.Workbook()
    out.remove(out.active)
    dorsales = sorted(plantilla)
    for (anio, mes), mcs in sorted(meses.items()):
        ws = out.create_sheet(f"{MESES[mes]} {anio}")
        _hoja_mes(ws, f"{MESES[mes].upper()} {anio}", mcs, dorsales, plantilla)
    out.save(ruta)
    return ruta


def _hoja_mes(ws, titulo, mcs, dorsales, plantilla):
    cols = [ev for _, evs in mcs for ev in evs]
    ncol = 3 + len(cols) + 2
    ws.cell(1, 1, f"DISPONIBILIDAD Y MINUTOS — {titulo}").font = Font(bold=True, size=13, color="FF1F3864")
    ws.cell(2, 1, "Minutos por sesión/partido. Blanco = disponible · Amarillo = Rehab (trabajo "
                  "individual, sin objetivo) · Rojo = lesionado / no disponible · 0 sin color = "
                  "descanso, gestión de carga o convocado sin jugar. El microciclo completo va al "
                  f"mes en que empieza. {MARCA} a partir de las hojas de microciclo.")
    ws.cell(2, 1).alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[2].height = 42
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncol)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncol)

    for j, h in enumerate(["Dorsal", "Jugador", "Grupo"], start=1):
        c = ws.cell(4, j, h)
        c.fill, c.font = AZUL_HDR, BLANCA
    col = 4
    inicios = []
    for etiqueta, evs in mcs:
        inicios.append(col)
        ws.cell(3, col, etiqueta).fill = MC_HDR
        ws.cell(3, col).font = BLANCA
        ws.cell(3, col).alignment = Alignment(horizontal="center")
        if len(evs) > 1:
            ws.merge_cells(start_row=3, start_column=col, end_row=3, end_column=col + len(evs) - 1)
        for key, fecha, partido, _ in evs:
            c = ws.cell(4, col, f"{key}\n{fecha:%d/%m}" + (" (MD)" if partido else ""))
            c.fill, c.font = (ORO_HDR if partido else AZUL_HDR), BLANCA
            c.alignment = Alignment(wrap_text=True, horizontal="center")
            col += 1
    c_tot, c_pct = col, col + 1
    for c, h in ((c_tot, "TOTAL MIN."), (c_pct, "% DISPON.")):
        ws.cell(4, c, h).fill, ws.cell(4, c).font = AZUL_HDR, BLANCA
    ws.row_dimensions[4].height = 30

    totales = {}
    for i, d in enumerate(dorsales):
        r = 5 + i
        nombre, grupo = plantilla[d]
        banda = BANDA[i % 2]
        for j, v in enumerate([d, nombre, grupo], start=1):
            ws.cell(r, j, v).fill = banda
        tot = 0
        for k, (_, _, _, res) in enumerate(cols):
            mins, est = res.get(d, (0, None))
            c = ws.cell(r, 4 + k, mins)
            c.fill = ROJO if est == C.LESION else AMARILLO if est == C.REHAB else banda
            c.number_format = "0.0"
            tot += mins or 0
        totales[d] = round(tot, 1)
        ws.cell(r, c_tot, totales[d]).fill = banda
        ws.cell(r, c_tot).font = Font(bold=True)
    fila_max = 5 + len(dorsales)
    max_tot = max(totales.values()) if totales else 0
    for i, d in enumerate(dorsales):
        c = ws.cell(5 + i, c_pct, round(totales[d] / max_tot, 4) if max_tot else 0)
        c.number_format, c.fill = "0%", BANDA[i % 2]
    ws.cell(fila_max, 2, "MÁXIMO EQUIPO").font = Font(bold=True)
    for j in range(1, ncol + 1):
        ws.cell(fila_max, j).fill = GRIS
    for k in range(len(cols)):
        vals = [num(ws.cell(5 + i, 4 + k).value) or 0 for i in range(len(dorsales))]
        ws.cell(fila_max, 4 + k, max(vals) if vals else 0).number_format = "0.0"
    ws.cell(fila_max, c_tot, max_tot)
    # línea gruesa al empezar cada microciclo
    for c0 in inicios:
        for r in range(3, fila_max + 1):
            cell = ws.cell(r, c0)
            cell.border = Border(left=GRUESA)
    ws.column_dimensions["A"].width = 7
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 7
    for c in range(4, ncol + 1):
        ws.column_dimensions[get_column_letter(c)].width = 9
    ws.column_dimensions[get_column_letter(c_tot)].width = 11
    ws.freeze_panes = "D5"
