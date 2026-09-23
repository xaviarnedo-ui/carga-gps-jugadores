"""Utilidades openpyxl comunes a todas las hojas."""
import os
import re
import shutil
import datetime as dt
from copy import copy

from openpyxl.styles import Font, PatternFill

from . import config as C

SIN_RELLENO = PatternFill(fill_type=None)
BLANCO = PatternFill("solid", fgColor="FFFFFFFF")


def vaciar(cell):
    """OJO bug openpyxl: ws.cell(r, c, value=None) NO borra. Siempre asignar el atributo."""
    cell.value = None


def pinta(cell, par):
    """par = (fill, font) de config (AZUL/VERDE/...) o None para quitar color."""
    font = copy(cell.font) if cell.font else Font()
    if par is None:
        cell.fill = SIN_RELLENO
        font.color = "FF000000"
    else:
        cell.fill = PatternFill("solid", fgColor=par[0])
        font.color = par[1]
    cell.font = font


def semaforo(real, obj):
    """Semáforo de cumplimiento. Objetivo 0 -> verde (no hay nada que cumplir ni pasarse)."""
    if real is None or obj is None:
        return None
    if obj == 0:
        return C.VERDE
    pct = (real - obj) / obj * 100
    if pct < -10:
        return C.AZUL
    if pct <= 10:
        return C.VERDE
    if pct <= 20:
        return C.NARANJA
    return C.ROJO


def semaforo_acwr(v):
    if v is None:
        return None
    if v < 0.80:
        return C.AZUL
    if v <= 1.30:
        return C.VERDE
    if v <= 1.50:
        return C.AMARILLO
    return C.ROJO


def fecha_hoja(ws):
    m = re.search(r"(\d{2})/(\d{2})/(\d{4})", str(ws.cell(1, 1).value or ""))
    return dt.date(int(m.group(3)), int(m.group(2)), int(m.group(1))) if m else None


def fila_cabecera(ws, clave="Dorsal", max_fila=10):
    for r in range(1, max_fila + 1):
        if ws.cell(r, 1).value == clave:
            return r
    raise ValueError(f"{ws.title}: no encuentro la fila de cabecera '{clave}'")


def filas_jugadores(ws):
    """-> ({dorsal: fila}, fila_media). Lee desde la cabecera hasta 'MEDIA EQUIPO'."""
    r0 = fila_cabecera(ws) + 1
    filas, media = {}, None
    for r in range(r0, r0 + 40):
        a, b = ws.cell(r, 1).value, ws.cell(r, 2).value
        if isinstance(a, (int, float)) and not isinstance(a, bool):
            filas[int(a)] = r
        elif str(b or "").strip().upper().startswith("MEDIA"):
            media = r
            break
        elif a is None and b is None and filas:
            media = r            # CARGA_AC: la fila de media a veces va sin etiqueta
            break
    return filas, media


def plantilla_de(ws):
    filas, _ = filas_jugadores(ws)
    return {d: ws.cell(r, 2).value for d, r in filas.items()}, \
           {d: ws.cell(r, 3).value for d, r in filas.items()}


def num(v):
    return v if isinstance(v, (int, float)) and not isinstance(v, bool) else None


def r1(x):
    return None if x is None else round(x + 0.0, 1)


def limpio(x):
    """Número bonito para Excel: 4504.0 -> 4504, 3463.44 -> 3463.4 (ya redondeado fuera)."""
    if x is None:
        return None
    return int(x) if float(x) == int(x) else x


def backup(rutas):
    """Copia los ficheros a GPS/_backups/<timestamp>/ antes de tocarlos."""
    stamp = dt.datetime.now().strftime("%Y%m%d-%H%M%S")
    dest = os.path.join(C.BACKUP_DIR, stamp)
    os.makedirs(dest, exist_ok=True)
    for r in rutas:
        if os.path.exists(r):
            shutil.copy2(r, dest)
    return dest


def ruta_microciclo(n):
    return os.path.join(C.MICRO_DIR, f"Microciclo {n}.xlsx")


def microciclos_existentes():
    out = {}
    for f in os.listdir(C.MICRO_DIR):
        m = re.match(r"^Microciclo (\d+)\.xlsx$", f)
        if m:
            out[int(m.group(1))] = os.path.join(C.MICRO_DIR, f)
    return dict(sorted(out.items()))
