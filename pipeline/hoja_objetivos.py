"""Hoja PDF de objetivos de una sesión para llevar al entrenamiento (A4 apaisado).

Objetivo de cada jugador por métrica + zona verde (±10%), agrupado por posición; quien no tiene
objetivo (lesión, rehab, descanso) va aparte. HTML → PDF con Chrome headless.
"""
import datetime as dt
import os
import subprocess
import tempfile

import openpyxl

from . import config as C
from . import estados as E
from . import sesion
from .xlsx import fecha_hoja, filas_jugadores, microciclos_existentes, num

CHROME = "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome"
SALIDA = os.path.join(C.GPS_DIR, "Informes", "Objetivos")
DIAS = ["Lunes", "Martes", "Miércoles", "Jueves", "Viernes", "Sábado", "Domingo"]
POS = [("D", "Defensas"), ("M", "Medios"), ("DL", "Delanteros")]
COLS = [("distancia", "Distancia", "m", 0), ("hmld", "HMLD", "m", 0), ("hsr", "HSR", "m", 0),
        ("sprint", "Sprints", "", 1), ("acc", "ACC", "", 1), ("dec", "DEC", "", 1)]
ESTADO_TXT = {C.LESION: "Lesión", C.REHAB: "Readaptación", C.DESCANSO: "Descanso"}


def _fmt(v, dec):
    if v is None:
        return "–"
    s = f"{v:,.{dec}f}".replace(",", "X").replace(".", ",").replace("X", ".")
    return s[:-2] if dec and s.endswith(",0") else s


def _buscar(key):
    for n, ruta in sorted(microciclos_existentes().items()):
        wb = openpyxl.load_workbook(ruta)
        if f"{key}_GPS" in wb.sheetnames:
            return n, wb, wb[f"{key}_GPS"]
    raise ValueError(f"No encuentro la sesión {key}")


def generar(key):
    n, wb, ws = _buscar(key)
    fecha = fecha_hoja(ws)
    tipo, dia = sesion.tipo_y_dia(ws)
    filas, media = filas_jugadores(ws)
    est = E.cargar()
    fuera = E.fuera_media(fecha, est)
    partido = next((s for s in wb.sheetnames if s.startswith("J") and s.endswith("_GPS")), None)
    rival = ""
    if partido:
        a1 = str(wb[partido].cell(1, 1).value or "")
        rival = a1.split("vs ", 1)[1].split("(")[0].strip() if "vs " in a1 else ""

    con, sin = {g: [] for g, _ in POS}, []
    for d, r in sorted(filas.items()):
        nombre, grupo = ws.cell(r, 2).value, (ws.cell(r, 3).value or "").strip()
        obj = {k: num(ws.cell(r, C.SES_OBJ_COL[k]).value) for k, *_ in COLS}
        if any(v is not None for v in obj.values()):
            con.setdefault(grupo, []).append((d, nombre, obj))
        else:
            sin.append((d, nombre, ESTADO_TXT.get(E.vigente(est, d), "Sin objetivo")))
    med = {k: num(ws.cell(media, C.SES_OBJ_COL[k]).value) for k, *_ in COLS}

    def celda(v, dec):
        if v is None:
            return '<td class="v">–</td>'
        return (f'<td class="v"><b>{_fmt(v, dec)}</b>'
                f'<span>{_fmt(v * 0.9, dec)}–{_fmt(v * 1.1, dec)}</span></td>')

    filas_html = []
    for g, titulo in POS:
        if not con.get(g):
            continue
        filas_html.append(f'<tr class="grp"><td colspan="{len(COLS) + 3}">{titulo}</td></tr>')
        for d, nombre, obj in con[g]:
            tag = ' <em>fuera de media</em>' if d in fuera else ""
            filas_html.append(f'<tr><td class="d">{d}</td><td class="n">{nombre}{tag}</td>'
                              + "".join(celda(obj[k], dec) for k, _, _, dec in COLS) + '<td class="nt"></td></tr>')
    filas_html.append('<tr class="med"><td></td><td class="n">Media equipo</td>'
                      + "".join(celda(med[k], dec) for k, _, _, dec in COLS) + '<td class="nt"></td></tr>')

    sin_html = " · ".join(f"<b>{d}</b> {nom} <span>({e})</span>" for d, nom, e in sin) or "—"
    notas_fm = "; ".join(sorted({v for v in fuera.values() if v}))
    html = f"""<!doctype html><html lang="es"><head><meta charset="utf-8"><title>Objetivos {key}</title><style>
@page {{ size: A4 landscape; margin: 7mm 10mm; }}
* {{ box-sizing: border-box; }}
body {{ font-family: -apple-system, Helvetica, Arial, sans-serif; color: #111; margin: 0; font-size: 10pt; }}
.head {{ display: flex; justify-content: space-between; align-items: flex-end; border-bottom: 2.5px solid #0e2a52; padding-bottom: 5px; margin-bottom: 6px; }}
h1 {{ font-size: 20pt; margin: 0; letter-spacing: -.01em; }} h1 small {{ font-size: 12pt; color: #555; font-weight: 600; }}
.club {{ text-align: right; font-size: 9pt; color: #444; line-height: 1.35; }}
table {{ width: 100%; border-collapse: collapse; }}
th {{ font-size: 9pt; text-transform: uppercase; letter-spacing: .04em; color: #444; border-bottom: 1.5px solid #333; padding: 3px 4px; text-align: center; }}
th small {{ text-transform: none; color: #777; font-weight: 400; }}
td {{ border-bottom: 1px solid #d8d8d8; padding: 1px 4px; height: 22px; }}
td.d {{ width: 26px; text-align: right; font-weight: 700; color: #555; }}
td.n {{ width: 175px; font-weight: 600; white-space: nowrap; }}
td.n em {{ font-style: normal; font-size: 7.5pt; font-weight: 700; color: #8a6d16; background: #fbf1d0; border-radius: 8px; padding: 0 5px; margin-left: 3px; }}
td.v {{ text-align: center; font-variant-numeric: tabular-nums; }}
td.v b {{ font-size: 11.5pt; display: block; line-height: 1; }}
td.v span {{ font-size: 7pt; color: #2b8a3e; display: block; line-height: 1.05; }}
td.nt {{ width: 120px; border-left: 1px solid #d8d8d8; }}
tr.grp td {{ font-size: 8pt; font-weight: 700; text-transform: uppercase; letter-spacing: .06em; color: #0e2a52; background: #eef1f6; height: 15px; padding: 0 6px; }}
tr.med td {{ border-top: 1.5px solid #333; background: #f6f6f3; }}
tr:nth-child(even):not(.grp):not(.med) td {{ background: #fafafa; }}
.foot {{ margin-top: 6px; font-size: 8.5pt; color: #333; line-height: 1.45; }}
.foot span {{ color: #777; }}
</style></head><body>
<div class="head"><h1>{key} · {dia} <small>Tipo {tipo} · {DIAS[fecha.weekday()]} {fecha:%d/%m/%Y}</small></h1>
<div class="club">AT BALEARES · Microciclo {n}{(" · " + partido[:-4] + " vs " + rival) if rival else ""}<br>
Objetivo individual = REF_PARTIDO × coeficiente {dia} (Tipo {tipo}) · en verde, zona ±10%</div></div>
<table><thead><tr><th>#</th><th style="text-align:left">Jugador</th>"""
    html += "".join(f"<th>{t}{(' <small>(' + u + ')</small>') if u else ''}</th>" for _, t, u, _ in COLS)
    html += f"""<th>Notas</th></tr></thead><tbody>{''.join(filas_html)}</tbody></table>
<div class="foot"><b>Sin objetivo hoy:</b> {sin_html}"""
    if fuera:
        html += f"<br><b>Fuera de media</b> (objetivo individual, no cuentan para la media del equipo){': ' + notas_fm if notas_fm else ''}."
    html += f"""<br><span>Generado el {dt.date.today():%d/%m/%Y} desde el Excel del Microciclo {n}.</span></div>
</body></html>"""

    os.makedirs(SALIDA, exist_ok=True)
    pdf = os.path.join(SALIDA, f"Objetivos {key} {fecha:%d-%m}.pdf")
    with tempfile.NamedTemporaryFile("w", suffix=".html", delete=False, encoding="utf-8") as f:
        f.write(html)
    subprocess.run([CHROME, "--headless=new", "--disable-gpu", "--no-pdf-header-footer",
                    f"--print-to-pdf={pdf}", "file://" + f.name], check=True, capture_output=True)
    os.unlink(f.name)
    return pdf
