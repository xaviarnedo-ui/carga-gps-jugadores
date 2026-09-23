"""Apertura de un microciclo nuevo a partir del anterior (estructura fija Tipo A/B/C):

Lunes MD+1 · Martes descanso · Miércoles MD-4 · Jueves MD-3 · Viernes MD-2 · Sábado MD-1 ·
Domingo partido. Objetivo = REF_PARTIDO vigente × coeficiente del día; MD+1 según el rol en el
último partido (titular ≥60', suplente <60' o no jugó); lesión/rehab = sin objetivo.
"""
import datetime as dt
import os
import re

import openpyxl

from . import acumulado, carga_ac, estados as E, referencia, sesion
from . import config as C
from .xlsx import (filas_jugadores, fecha_hoja, microciclos_existentes, num, pinta,
                   ruta_microciclo, vaciar)

DIAS_SEMANA = ["Lunes", "Martes", "Miércoles", "Jueves", "Viernes", "Sábado", "Domingo"]
MESES_CORTOS = ["", "ENE", "FEB", "MAR", "ABR", "MAY", "JUN", "JUL", "AGO", "SEP", "OCT", "NOV", "DIC"]
ESTRUCTURA = [(0, "MD+1"), (2, "MD-4"), (3, "MD-3"), (4, "MD-2"), (5, "MD-1")]   # (offset, día)
TITULAR_MIN = 60
LEYENDA = ("LEYENDA: AZUL < −10% (corto) · VERDE ±10% (cumplido) · NARANJA +10% a +20% (pasado) "
           "· ROJO > +20% (muy pasado).")


def roles_md1(ws_partido):
    """{dorsal: 'T'|'S'} según los minutos del último partido."""
    filas, _ = filas_jugadores(ws_partido)
    out = {}
    for d, r in filas.items():
        v = ws_partido.cell(r, C.MAT_DUR).value
        mins = None
        if v not in (None, ""):
            from .pdf_catapult import dur_a_min
            mins = dur_a_min(v)
        out[d] = "T" if mins is not None and mins >= TITULAR_MIN else "S"
    return out


def _limpia_filas(ws, cols, desde_fila, hasta_fila):
    for r in range(desde_fila, hasta_fila + 1):
        for c in cols:
            vaciar(ws.cell(r, c))
            pinta(ws.cell(r, c), None)


def abrir(tipo, partido_key, rival, lunes, estados=None, ref=None, coefs=None,
          destino=None, n_prev=None, liga=True):
    existentes = microciclos_existentes()
    n_prev = n_prev or max(existentes)
    n = n_prev + 1
    destino = destino or ruta_microciclo(n)
    if os.path.exists(destino):
        raise FileExistsError(f"Ya existe {destino}")
    if lunes.weekday() != 0:
        raise ValueError(f"{lunes} no es lunes")

    wb = openpyxl.load_workbook(existentes[n_prev])
    estados = estados if estados is not None else E.cargar()
    if ref is None or coefs is None:
        wt = openpyxl.load_workbook(C.TIPO_XLSX)
        ref = ref or referencia.ref_partido(wt)
        coefs = coefs or referencia.coeficientes(wt)
    if tipo not in coefs:
        raise ValueError(f"No hay coeficientes para el Tipo {tipo} en Microciclo_Tipo.xlsx")

    ses = sesion.hojas_sesion(wb)
    partidos = [wb[x] for x in wb.sheetnames if re.match(r"^(PT|J)\d+_GPS$", x)]
    if len(ses) != 5 or len(partidos) != 1:
        raise ValueError(f"Microciclo {n_prev}: esperaba 5 sesiones y 1 partido para copiar la "
                         f"estructura; hay {len(ses)} y {len(partidos)}")
    ws_part_prev = partidos[0]
    prev_key = ws_part_prev.title[:-4]
    roles = roles_md1(ws_part_prev)
    for x in [x for x in wb.sheetnames if x.lower().startswith("extra")]:
        del wb[x]

    ult = max(int(re.match(r"S(\d+)", k).group(1)) for k, _, _ in ses)
    claves = []
    for i, ((_, ws, _), (off, dia)) in enumerate(zip(ses, ESTRUCTURA)):
        key = f"S{ult + 1 + i}"
        fecha = lunes + dt.timedelta(days=off)
        claves.append((key, dia, fecha))
        ws.title = f"{key}_GPS"
        ws.cell(1, 1).value = (f"SESIÓN {key} ({DIAS_SEMANA[fecha.weekday()]} {fecha:%d/%m/%Y}) "
                               f"— OBJETIVO vs REAL")
        extra = (" o Suplentes (<60' o no jugó)" if dia == "MD+1" else "")
        coef_txt = (f"coeficiente +1 Titulares (jugó ≥60' en {prev_key}){extra}"
                    if dia == "MD+1" else f"coeficiente {dia} de Tipo {tipo}")
        ws.cell(2, 1).value = (f"Microciclo {n} — Tipo {tipo}, objetivo {dia}. Objetivo INDIVIDUAL: "
                               f"REF_PARTIDO propio de cada jugador × {coef_txt}. Real pendiente de "
                               f"cargar. {LEYENDA}")
        ws.cell(3, 1).value = f"MICROCICLO {n} · TIPO {tipo} · {dia}"
        filas, media = filas_jugadores(ws)
        cols = [c + k for c in C.SES_OBJ_COL.values() for k in (1, 2)] + [C.SES_VMAX, C.SES_PL, C.SES_DUR]
        _limpia_filas(ws, cols, min(filas.values()), media)
        for c in C.SES_OBJ_COL.values():
            vaciar(ws.cell(media, c))
        for d, r in filas.items():
            est = E.vigente(estados, d)
            if est != C.FULL or d not in ref:
                obj = None
            else:
                clave_dia = dia + roles.get(d, "S") if dia == "MD+1" else dia
                obj = referencia.objetivo(ref[d], coefs[tipo][clave_dia])
            sesion._escribe_obj(ws, r, obj)
        sesion.escribir_media(ws, filas, media, {d: C.FULL for d, r in filas.items()
                                                 if sesion.tiene_objetivo(ws, r)})
        for rr in range(media + 1, media + 6):
            vaciar(ws.cell(rr, 1))
        ws.cell(media + 1, 1).value = (
            f"Objetivo {dia} (Tipo {tipo}) — INDIVIDUAL por jugador (REF_PARTIDO propio × coeficiente"
            + (f" Titular/Suplente según minutos jugados en {prev_key})." if dia == "MD+1" else ")."))
        nombres = {d: ws.cell(r, 2).value for d, r in filas.items()}
        est_dia = {d: E.vigente(estados, d) for d in filas}
        ws.cell(media + 2, 1).value = sesion.nota_estados(nombres, est_dia)

    # partido
    fecha_p = lunes + dt.timedelta(days=6)
    ws = ws_part_prev
    ws.title = f"{partido_key}_GPS"
    jornada = re.match(r"J(\d+)", partido_key)
    comp = f"Liga, Jornada {jornada.group(1)}" if jornada else "Amistoso"
    ws.cell(1, 1).value = (f"PARTIDO {partido_key} · vs {rival} ({fecha_p:%d/%m/%Y}) — DATOS DE PARTIDO "
                           f"({comp})")
    ws.cell(2, 1).value = "No lleva Obj/Dif. Datos pendientes de cargar tras el encuentro."
    ws.cell(3, 1).value = f"MICROCICLO {n} · PARTIDO {'LIGA ' if jornada else ''}{partido_key}"
    filas, media = filas_jugadores(ws)
    _limpia_filas(ws, range(4, 13), min(filas.values()), media)
    for rr in range(media + 1, media + 8):
        vaciar(ws.cell(rr, 1))

    # Acumulado
    wa = wb["Acumulado"]
    fin = lunes + dt.timedelta(days=6)
    rango = (f"{lunes.day}-{fin.day} {MESES_CORTOS[fin.month]}" if lunes.month == fin.month
             else f"{lunes.day} {MESES_CORTOS[lunes.month]}-{fin.day} {MESES_CORTOS[fin.month]}")
    partes = [f"{k} ({dia})" for k, dia, _ in claves]
    partes.insert(1, "Descanso (Martes)")
    wa.cell(3, 1).value = (f"SEMANA: MICROCICLO {n} ({rango}) · TIPO {tipo} · SESIONES: "
                           + " · ".join(partes)
                           + f" · {partido_key} vs {rival} ({comp.split(',')[0]}, {fecha_p:%d/%m})")
    wa.cell(4, 1).value = (f"Objetivo de la semana completa de Microciclo {n} (Lunes MD+1 individual "
                           f"según rol en {prev_key} + Miércoles MD-4 + Jueves MD-3 + Viernes MD-2 + "
                           f"Sábado MD-1). Acumulado (Real) pendiente de cargar sesión a sesión.")
    filas_a, media_a = filas_jugadores(wa)
    for rr in range(media_a + 1, media_a + 5):
        vaciar(wa.cell(rr, 1))
    wa.cell(media_a + 2, 1).value = f"Microciclo {n} EN CURSO."
    acumulado.recalcular(wb)

    # CARGA_AC: cabeceras nuevas y columnas vacías
    wc = wb["CARGA_AC"]
    hdr = [r for r in range(1, 10) if wc.cell(r, 1).value == "Dorsal"][0]
    eventos = [(k, f) for k, _, f in claves] + [(partido_key, fecha_p)]
    for b in carga_ac._bloques(wc, hdr):
        if len(b["dias"]) != len(eventos):
            raise ValueError("CARGA_AC: el número de columnas por sesión no cuadra con la semana")
        pref = {"pl": "PL", "hsr": "HSR", "sprint": "Sprint"}[b["metrica"]]
        for (col, _), (k, f) in zip(b["dias"], eventos):
            suf = f" ({comp.split(',')[0]})" if k == partido_key else ""
            wc.cell(hdr, col).value = f"{pref} {k}\n{f:%d/%m}{suf}"
    wb.save(destino)

    # ACWR de arranque: a fecha del último dato disponible (el nuevo aún no tiene ninguno)
    wb = openpyxl.load_workbook(destino)
    hist = carga_ac.historial({n: wb} if os.path.dirname(destino) == C.MICRO_DIR else None)
    ultimo = max((f for serie in hist.values() for (_, f) in serie), default=None)
    carga_ac.rellenar(wb, n, hist, asof=ultimo)
    wb.save(destino)
    return destino, claves, (partido_key, fecha_p), roles
