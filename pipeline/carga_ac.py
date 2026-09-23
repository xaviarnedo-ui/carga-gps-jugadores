"""CARGA_AC: histórico diario de PL/HSR/Sprint de TODOS los microciclos → ACWR.

aguda = Σ últimos 7 días naturales / 7 · crónica = Σ últimos 28 / 28 · ACWR = aguda / crónica
(días sin sesión = 0; crónica 0 → ACWR 0). La media de equipo excluye ACWR = 0.
Fuente: las hojas *_GPS (sesiones, partidos y Extra) — no las columnas de CARGA_AC.
"""
import datetime as dt
import re
from collections import defaultdict

import openpyxl

from . import config as C
from . import sesion
from .xlsx import (fila_cabecera, filas_jugadores, fecha_hoja, limpio, microciclos_existentes,
                   num, pinta, r1, semaforo_acwr, vaciar)

CLAVES = {"PL": "pl", "HSR": "hsr", "SPRINT": "sprint"}


def _cols(nombre_hoja):
    if sesion.es_hoja_sesion(nombre_hoja):
        return {"pl": C.SES_PL, "hsr": C.SES_OBJ_COL["hsr"] + 1, "sprint": C.SES_OBJ_COL["sprint"] + 1}
    return {"pl": C.MAT_PL, "hsr": C.MAT_COL["hsr"], "sprint": C.MAT_COL["sprint"]}


def hojas_gps(wb):
    return [wb[n] for n in wb.sheetnames if n.endswith("_GPS")]


def valor_en_hoja(ws, dorsal, metrica):
    filas, _ = filas_jugadores(ws)
    if dorsal not in filas:
        return None
    return num(ws.cell(filas[dorsal], _cols(ws.title)[metrica]).value)


def historial(abiertos=None):
    """{métrica: {(dorsal, fecha): valor}} sumando todas las hojas de todos los microciclos.
    abiertos = {n: wb} para usar los libros en memoria (con cambios aún sin guardar)."""
    abiertos = abiertos or {}
    hist = {m: defaultdict(float) for m in CLAVES.values()}
    for n, ruta in microciclos_existentes().items():
        wb = abiertos.get(n) or openpyxl.load_workbook(ruta, data_only=True)
        for ws in hojas_gps(wb):
            fecha = fecha_hoja(ws)
            if not fecha:
                continue
            filas, _ = filas_jugadores(ws)
            cols = _cols(ws.title)
            for dorsal, fila in filas.items():
                for m, c in cols.items():
                    v = num(ws.cell(fila, c).value)
                    if v is not None:
                        hist[m][(dorsal, fecha)] += v
    return hist


def acwr(serie, dorsal, asof):
    a = sum(serie.get((dorsal, asof - dt.timedelta(i)), 0) for i in range(7)) / 7
    c = sum(serie.get((dorsal, asof - dt.timedelta(i)), 0) for i in range(28)) / 28
    return (a / c if c else 0.0), a, c


def fecha_calculo(wb):
    """Último día con datos en este microciclo (sesión, partido o Extra)."""
    fechas = []
    for ws in hojas_gps(wb):
        filas, _ = filas_jugadores(ws)
        pl = _cols(ws.title)["pl"]
        if any(num(ws.cell(r, pl).value) is not None for r in filas.values()):
            fechas.append(fecha_hoja(ws))
    return max(fechas) if fechas else None


def _bloques(ws, fila_hdr):
    """-> [{'metrica', 'acwr': col, 'aguda': col, 'cronica': col, 'dias': [(col, key)]}]"""
    bloques = []
    for c in range(1, ws.max_column + 1):
        h = str(ws.cell(fila_hdr, c).value or "")
        m = re.match(r"\s*ACWR\s*(PL|HSR|Sprint)?\s*$", h, re.I)
        if m:
            bloques.append({"metrica": CLAVES[(m.group(1) or "PL").upper()], "acwr": c, "dias": []})
            continue
        if not bloques:
            continue
        b = bloques[-1]
        if re.match(r"(PL|HSR|Sprint)\s+(\S+)\n", h, re.I):
            b["dias"].append((c, re.match(r"\S+\s+(\S+)", h).group(1)))
        elif "aguda" in h.lower():
            b["aguda"] = c
        elif "crónica" in h.lower() or "cronica" in h.lower():
            b["cronica"] = c
    return bloques


def rellenar(wb, n, hist, asof=None):
    """Columnas por sesión (0 para quien no tuvo sesión) + ACWR/aguda/crónica a `asof`."""
    ws = wb["CARGA_AC"]
    hdr = fila_cabecera(ws)
    filas, fila_media = filas_jugadores(ws)
    asof = asof or fecha_calculo(wb)
    salida = {}
    for b in _bloques(ws, hdr):
        met = b["metrica"]
        for col, key in b["dias"]:
            hoja = wb[f"{key}_GPS"] if f"{key}_GPS" in wb.sheetnames else None
            cargada = hoja is not None and fecha_calculo_hoja(hoja)
            for dorsal, fila in filas.items():
                cell = ws.cell(fila, col)
                if not cargada:
                    vaciar(cell)
                else:
                    v = valor_en_hoja(hoja, dorsal, met)
                    cell.value = v if v is not None else 0
        activos = []
        for dorsal, fila in filas.items():
            if asof is None:
                break
            ratio, a, c = acwr(hist[met], dorsal, asof)
            ws.cell(fila, b["acwr"]).value = round(ratio, 2)
            ws.cell(fila, b["aguda"]).value = limpio(r1(a))
            ws.cell(fila, b["cronica"]).value = limpio(r1(c))
            pinta(ws.cell(fila, b["acwr"]), semaforo_acwr(round(ratio, 2)))
            salida.setdefault(dorsal, {})[met] = round(ratio, 2)
            if round(ratio, 2) > 0:
                activos.append((ratio, a, c))
        if fila_media and activos:
            k = len(activos)
            media = round(sum(x[0] for x in activos) / k, 2)
            ws.cell(fila_media, b["acwr"]).value = media
            ws.cell(fila_media, b["aguda"]).value = limpio(r1(sum(x[1] for x in activos) / k))
            ws.cell(fila_media, b["cronica"]).value = limpio(r1(sum(x[2] for x in activos) / k))
            pinta(ws.cell(fila_media, b["acwr"]), semaforo_acwr(media))

    if asof:
        a2 = ws.cell(2, 1)
        estado = "CERRADO" if microciclo_completo(wb) else "EN CURSO"
        txt = re.sub(r"\s*Microciclo \d+ (EN CURSO|CERRADO)[^.]*\.?", "", str(a2.value or "")).strip()
        a2.value = (f"{txt} Microciclo {n} {estado} (fecha de cálculo {asof:%d/%m/%Y})."
                    ).strip()
    return asof, salida


def fecha_calculo_hoja(ws):
    filas, _ = filas_jugadores(ws)
    pl = _cols(ws.title)["pl"]
    return any(num(ws.cell(r, pl).value) is not None for r in filas.values())


def microciclo_completo(wb):
    hojas = [ws for ws in hojas_gps(wb) if C.SHEET_KEY_RE.match(ws.title)]
    return bool(hojas) and all(fecha_calculo_hoja(ws) for ws in hojas)
