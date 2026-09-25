"""Regresión de hojas de partido: J3 (con fallo de GPS de Hernández) y J2."""
import os
import sys
import unittest

import openpyxl

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from pipeline import config as C, informe, partido, referencia  # noqa: E402
from pipeline.pdf_catapult import dur_a_min  # noqa: E402
from pipeline.xlsx import filas_jugadores, plantilla_de, vaciar  # noqa: E402

BASE = os.path.dirname(C.GPS_DIR)
PARTIDOS = os.path.join(C.GPS_DIR, "Partidos")
# REF_PARTIDO de Hernández (14) con la que se reconstruyó su J3 (fallo de GPS, 50'), antes del
# recálculo del 25/09 que pasó a extrapolar a la duración real de cada partido
REF_ANTES_J3 = {14: {"distancia": 9750.2, "hmld": 1526.7, "hsr": 210.2, "sprint": 2,
                     "acc": 35.2, "dec": 50.2}}


def snapshot(ws, ignorar=()):
    filas, media = filas_jugadores(ws)
    out = {}
    for r in list(filas.values()) + [media]:
        for c in range(4, 13):
            if (r, c) in ignorar:
                continue
            v = ws.cell(r, c).value
            if c == C.MAT_DUR and v:
                v = dur_a_min(v)                       # J2 guardaba H:MM:SS
            if r == media and isinstance(v, float) and c != C.MAT_VMAX:
                v = round(v, 1)                        # J2 guardaba la media con 2 decimales
            out[(r, c)] = v
    return out


class RegresionPartido(unittest.TestCase):
    def caso(self, n, key, pdf, fallos=None):
        ruta = os.path.join(C.MICRO_DIR, f"Microciclo {n}.xlsx")
        if not (os.path.exists(ruta) and os.path.exists(pdf)):
            self.skipTest("faltan ficheros del club")
        wb = openpyxl.load_workbook(ruta)
        ws = wb[f"{key}_GPS"]
        filas, media = filas_jugadores(ws)
        # Vel. máx / PL del jugador con GPS roto no son reproducibles (se pusieron a mano),
        # y por tanto tampoco la media de esas dos columnas
        ignorar = {(r, c) for r in [filas[d] for d in (fallos or {})] + ([media] if fallos else [])
                   for c in (C.MAT_VMAX, C.MAT_PL)}
        antes = snapshot(ws, ignorar)
        nombres, _ = plantilla_de(ws)
        datos, sin = informe.datos_por_dorsal(informe.leer(pdf), nombres)
        self.assertEqual(sin, [])
        ref_j3 = REF_ANTES_J3
        for d, mins in (fallos or {}).items():
            # J3 se reconstruyó con el criterio antiguo (extrapolación a 95')
            datos[d] = partido.estimar_por_fallo(ref_j3[d], mins, 95, partido.pl_por_metro(datos))
        for r in filas.values():
            for c in range(4, 13):
                vaciar(ws.cell(r, c))
        partido.rellenar(ws, datos)
        despues = snapshot(ws, ignorar)
        difs = [(k, antes[k], despues[k]) for k in antes if antes[k] != despues[k]]
        self.assertEqual(difs, [], difs[:10])

    def test_j3_con_fallo_gps(self):
        self.caso(10, "J3", os.path.join(PARTIDOS, "J3 UCAM.pdf"), fallos={14: 50})

    def test_j2(self):
        self.caso(9, "J2", os.path.join(PARTIDOS, "J2 INTERCITY.pdf"))


class Formulas(unittest.TestCase):
    def test_fatiga(self):
        self.assertEqual(referencia.estimar(1000, 95, 95), 1000)
        self.assertEqual(referencia.estimar(1000, 100.6, 100.6), 1000)
        self.assertAlmostEqual(referencia.estimar(5000, 50, 95), 5000 + 45 * 100 * 0.9)
        self.assertAlmostEqual(referencia.estimar(5000, 50, 100.6), 5000 + 50.6 * 100 * 0.9)

    def test_inversa_es_inversa(self):
        for mins in (20, 50, 80):
            est = referencia.estimar(4321, mins, 97.3)
            self.assertAlmostEqual(referencia.real_desde_ref(est, mins, 97.3), 4321)

    def test_duracion_partido(self):
        self.assertEqual(referencia.duracion_partido([50.1, 100.6, None, 81]), 100.6)


if __name__ == "__main__":
    unittest.main()
