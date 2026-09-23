"""Regresión: vaciar una sesión ya cargada, reprocesar su PDF y comparar con el Excel original.

Usa los Excel/PDF reales del club (fuera del repo); si no están, los tests se saltan.
"""
import os
import sys
import unittest

import openpyxl

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from pipeline import config as C, informe, sesion  # noqa: E402
from pipeline.pdf_catapult import dur_a_min  # noqa: E402
from pipeline.xlsx import filas_jugadores, num, plantilla_de, vaciar  # noqa: E402

BASE = os.path.dirname(C.GPS_DIR)
# (microciclo, sesión, pdf, comparar_colores, comparar_media)
# - Hasta M9 el Dif no se coloreaba y la duración iba como texto H:MM:SS.
# - La MEDIA EQUIPO original de S44 dejó fuera a Martín M. sin motivo anotado, y la de S38
#   tampoco sale con "todos los full": errores del proceso manual, aquí no se comparan.
CASOS = [
    (11, "S46", os.path.join(BASE, "S46.pdf"), True, True),
    (10, "S45", os.path.join(BASE, "S45.pdf"), True, True),
    (10, "S44", os.path.join(BASE, "S44.pdf"), True, False),
    (10, "S43", os.path.join(C.GPS_DIR, "Sesiones", "S43.pdf"), True, True),
    (9, "S38", os.path.join(C.GPS_DIR, "Sesiones", "S38.pdf"), False, False),
]


def estados_desde_hoja(ws):
    filas, _ = filas_jugadores(ws)
    out = {}
    for d, r in filas.items():
        obj = sesion.tiene_objetivo(ws, r)
        real = num(ws.cell(r, C.SES_PL).value) is not None or \
            any(num(ws.cell(r, c + 1).value) is not None for c in C.SES_OBJ_COL.values())
        out[d] = C.FULL if obj and real else C.REHAB if real else C.LESION
    return out


def snapshot(ws, colores, con_media):
    filas, media = filas_jugadores(ws)
    out = {}
    for r in list(filas.values()) + ([media] if con_media else []):
        for c in range(4, 25):
            cell = ws.cell(r, c)
            v = cell.value
            if c == C.SES_DUR:
                v = dur_a_min(v) if v else None
            color = cell.fill.fgColor.rgb if cell.fill.fill_type else None
            # la fila MEDIA EQUIPO a veces iba coloreada y a veces no: ahora siempre se colorea
            if not colores or c in C.SES_OBJ_COL.values() or r == media:
                color = None
            out[(r, c)] = (v, color)
    return out


class RegresionSesion(unittest.TestCase):
    def test_casos(self):
        for n, key, pdf, colores, con_media in CASOS:
            ruta = os.path.join(C.MICRO_DIR, f"Microciclo {n}.xlsx")
            if not (os.path.exists(ruta) and os.path.exists(pdf)):
                continue
            with self.subTest(sesion=key):
                wb = openpyxl.load_workbook(ruta)
                ws = wb[f"{key}_GPS"]
                antes = snapshot(ws, colores, con_media)
                estados = estados_desde_hoja(ws)
                nombres, _ = plantilla_de(ws)
                filas, _ = filas_jugadores(ws)
                # los datos de jugadores que no están en el PDF (rehab en informe aparte)
                # se toman de la propia hoja
                datos, sin = informe.datos_por_dorsal(informe.leer(pdf), nombres)
                self.assertEqual(sin, [])
                for d, e in estados.items():
                    if e in (C.FULL, C.REHAB) and d not in datos:
                        r = filas[d]
                        datos[d] = {m: ws.cell(r, c + 1).value for m, c in C.SES_OBJ_COL.items()}
                        datos[d].update(vmax=ws.cell(r, C.SES_VMAX).value,
                                        pl=ws.cell(r, C.SES_PL).value,
                                        dur=dur_a_min(ws.cell(r, C.SES_DUR).value))
                    if e == C.LESION and d in datos:
                        estados[d] = C.REHAB     # sin objetivo pero con dato -> rehab
                for r in filas.values():
                    for c in range(5, 25):
                        if (c - 4) % 3 != 0:
                            vaciar(ws.cell(r, c))
                sesion.rellenar(ws, datos, estados)
                despues = snapshot(ws, colores, con_media)
                difs = [(k, antes[k], despues[k]) for k in antes if antes[k] != despues[k]]
                self.assertEqual(difs, [], f"{key}: {len(difs)} celdas distintas: {difs[:8]}")


if __name__ == "__main__":
    unittest.main()
